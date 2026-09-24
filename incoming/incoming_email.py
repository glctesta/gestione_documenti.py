# -*- coding: utf-8 -*-
"""
incoming_email.py
Email + job schedulati del modulo "Ricezione" (Incoming).

Comprende:
  - Notifiche email (in thread daemon, mai propagano eccezioni):
      * nuova richiesta       -> send_new_request_email()
      * risposta pronta       -> send_answer_email()
      * escalation PENDING    -> process_escalations() (popup + email riepilogo)
      * reminder periodici    -> check_and_send_reminders()
  - Report mensile Excel    -> generate_and_send_monthly_report()
    (job 'incoming_monthly_report' con lock cross-PC su
    Traceability_RS.dbo.AutomaticEmailJobs via email_job_coordinator;
    destinatari da settings atribute='Incoming_soluzione_problemi').

Uso standalone (Task Scheduler / manuale):
  .venv\\Scripts\\python.exe incoming\\incoming_email.py reminders [--dry-run]
  .venv\\Scripts\\python.exe incoming\\incoming_email.py monthly [--year Y] [--month M]
                                                      [--force] [--dry-run]
"""
import logging
import os
import sys
import tempfile
import threading
from datetime import datetime

# Esecuzione diretta come script: aggiungi la project root a sys.path
# PRIMA dei moduli flat del progetto (kit_notifications, utils, ...).
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    from . import incoming_db
except ImportError:  # esecuzione come script standalone
    import incoming_db

from email_job_coordinator import (
    claim_job_run, force_claim_job, release_job_lock, log_job_run,
)
from kit_notifications import queue_popup
import utils

logger = logging.getLogger(__name__)

JOB_MONTHLY = 'incoming_monthly_report'
JOB_REMINDERS = 'incoming_reminders'

ESCALATION_AGE_MINUTES = 120
ESCALATION_INTERVAL_MINUTES = 30

# Job inseriti idempotentemente (WHERE NOT EXISTS) perche' non sono nella
# lista DEFAULT_JOBS di email_job_coordinator.py (file di altro agente).
_JOB_SPECS = {
    JOB_MONTHLY: (
        'Report mensile Ricezione (Incoming)',
        'generate_and_send_monthly_report',
        'incoming/incoming_email.py',
        "Report Excel mensile (e YTD) delle richieste incoming: per tipo, "
        "risposte, confermate OK/KO e tempo medio di risposta.",
        'Primo giorno lavorativo del mese 09:00',
        incoming_db.MONTHLY_RECIPIENTS_ATTRIBUTE,
    ),
    JOB_REMINDERS: (
        'Reminder richieste Ricezione PENDING',
        'check_and_send_reminders',
        'incoming/incoming_email.py',
        "Sollecito per le richieste incoming ancora PENDING, secondo la "
        "frequenza configurata per tipo (reminders_per_day).",
        'Ogni 30-60 min (Task Scheduler)',
        None,
    ),
}


def _cursor(db):
    """Cursor uniforme: funziona con la classe Database di main.py e con
    BackgroundDatabase del servizio background."""
    if hasattr(db, '_ensure_connection'):
        try:
            db._ensure_connection()
        except Exception:
            pass
    return db.cursor


def _ensure_job_row(db, job_name):
    """Inserisce la riga del job in AutomaticEmailJobs se mancante (idempotente)."""
    spec = _JOB_SPECS.get(job_name)
    if spec is None:
        return
    try:
        with db._lock:
            cur = _cursor(db)
            cur.execute(
                """
                INSERT INTO Traceability_RS.dbo.AutomaticEmailJobs
                    (JobName, DisplayName, FunctionName, ModulePath,
                     Description, Timing, RecipientsSettingKey, IsEnabled)
                SELECT ?, ?, ?, ?, ?, ?, ?, 1
                WHERE NOT EXISTS (
                    SELECT 1 FROM Traceability_RS.dbo.AutomaticEmailJobs
                    WHERE JobName = ?
                )
                """,
                (job_name, spec[0], spec[1], spec[2], spec[3], spec[4], spec[5], job_name)
            )
            db.conn.commit()
    except Exception as e:
        logger.warning("Impossibile creare job %s: %s", job_name, e)


# ───────────────────── Invio asincrono (mai propagano) ────────────────── #

def _send_email_async(recipients, subject, body, attachments=None):
    """Invia l'email in un thread daemon; errori solo loggati."""

    def _send():
        try:
            utils.send_email(recipients, subject, body, attachments=attachments)
        except Exception as e:
            logger.error("Invio email incoming fallito (%s): %s", subject, e)

    threading.Thread(target=_send, daemon=True).start()


def _recipients_for_type(db, request_type):
    try:
        cfg = incoming_db.get_email_config(db, request_type) or {}
        return list(cfg.get('emails') or [])
    except Exception as e:
        logger.error("Lettura destinatari per %s fallita: %s", request_type, e)
        return []


def _type_label(request_type):
    return incoming_db.REQUEST_TYPES.get(request_type, request_type or '-')


def _fmt_dt(value):
    try:
        return value.strftime('%d/%m/%Y %H:%M')
    except Exception:
        return str(value) if value else '-'


# ───────────────────── Template notifiche evento ──────────────────────── #

def new_request_messages(req: dict) -> dict:
    """Testi per la notifica di una nuova richiesta (email + popup gia'
    inserito dal chiamante con queue_popup)."""
    number = req.get('RequestNumber', '?')
    subject = f"[RICHIESTA] Ricezione — {_type_label(req.get('RequestType'))} — {number}"
    body = (
        f"Nuova richiesta incoming {number} da {req.get('RequestedBy', '?')} "
        f"(postazione {req.get('RequesterHost', '-')}):\n\n"
        f"Tipo: {_type_label(req.get('RequestType'))}\n"
        f"Fornitore: {req.get('SupplierName', '-')} (ID {req.get('SupplierId', req.get('Supplier_id', '-'))})\n"
        f"DDT: {req.get('DdtNumber', '-')} del {_fmt_dt(req.get('DdtDate'))}\n"
        f"P.O.: {req.get('PurOrderNumber', '-')}\n"
        f"MPN: {req.get('MpnCode', '-') or req.get('WrongMpn', '-')}\n"
        f"Codice interno: {req.get('ComponentCode') or '-'}\n"
        f"Quantita': {req.get('QtyToReceive', '-')} "
        f"(attesa da P.O.: {req.get('QtyExpectedPerPo', '-')})\n"
        f"Richiesta il: {_fmt_dt(req.get('RequestedOn'))}\n\n"
        f"Accedi al modulo Ricezione (Soluzioni Incoming) per rispondere."
    )
    return {'subject': subject, 'body': body}


def answer_messages(req: dict) -> dict:
    """Testi per la notifica di risposta pronta (il popup al richiedente
    e' gia' stato inserito dal chiamante con queue_popup)."""
    number = req.get('RequestNumber', '?')
    subject = f"[RISPOSTA] Ricezione — {number}"
    body = (
        f"La richiesta incoming {number} e' stata risolta da "
        f"{req.get('AnsweredBy', '?')}:\n\n"
        f"Tipo: {_type_label(req.get('RequestType'))}\n"
        f"MPN risposta: {req.get('AnswerMpnCode', '-')}\n"
        f"Risposta: {req.get('AnswerText', '-')}\n"
        f"Risposta il: {_fmt_dt(req.get('AnsweredOn'))}\n\n"
        f"Accedi al modulo Ricezione per confermare la soluzione."
    )
    return {'subject': subject, 'body': body}


# ───────────────────── Email evento (API per gli altri moduli) ────────── #

def send_new_request_email(db, req: dict) -> None:
    """Email di notifica nuova richiesta ai destinatari configurati per il tipo."""
    recipients = _recipients_for_type(db, req.get('RequestType'))
    if not recipients:
        logger.warning("Nessun destinatario per tipo %s: email nuova richiesta "
                       "%s NON inviata", req.get('RequestType'), req.get('RequestNumber'))
        return
    msg = new_request_messages(req)
    _send_email_async(recipients, msg['subject'], msg['body'])


def send_answer_email(db, req: dict) -> None:
    """Email di notifica risposta pronta ai destinatari configurati per il tipo."""
    recipients = _recipients_for_type(db, req.get('RequestType'))
    if not recipients:
        logger.warning("Nessun destinatario per tipo %s: email risposta %s "
                       "NON inviata", req.get('RequestType'), req.get('RequestNumber'))
        return
    msg = answer_messages(req)
    _send_email_async(recipients, msg['subject'], msg['body'])


# ───────────────────── Escalation ─────────────────────────────────────── #

def process_escalations(db, dry_run: bool = False) -> int:
    """Rileva le richieste PENDING in escalation (eta' >= 120 min, nessun
    popup di escalation da >= 30 min) e per ciascuna:
      - claim atomico su LastEscalationPopup (una sola postazione vince);
      - popup di escalation a 'INCOMING_RECEIVER' (category='INCOMING');
      - email riepilogativa ai destinatari configurati per il tipo.

    Ritorna il numero di richieste escalate. Sicuro su piu' PC concorrenti.
    """
    try:
        candidates = incoming_db.get_escalation_candidates(db)
    except Exception as e:
        logger.error("Lettura candidati escalation fallita: %s", e)
        return 0
    if not candidates:
        return 0

    claimed = []
    try:
        with db._lock:
            cur = _cursor(db)
            for c in candidates:
                if dry_run:
                    # Nessun claim: sola valutazione, nessuna modifica al DB
                    claimed.append(c)
                    continue
                # Claim atomico: la condizione ripete il filtro dei candidati,
                # cosi' una sola istanza (anche tra piu' PC receiver) vince.
                cur.execute(
                    """
                    UPDATE Traceability_RS.dyn.IncomingRequest
                    SET LastEscalationPopup = GETDATE()
                    WHERE Id = ? AND Status = ?
                      AND DATEDIFF(MINUTE, RequestedOn, GETDATE()) >= ?
                      AND (LastEscalationPopup IS NULL
                           OR DATEDIFF(MINUTE, LastEscalationPopup, GETDATE()) >= ?)
                    """,
                    (c.get('Id'), incoming_db.STATUS_PENDING,
                     ESCALATION_AGE_MINUTES, ESCALATION_INTERVAL_MINUTES)
                )
                if cur.rowcount > 0:
                    claimed.append(c)
            if claimed and not dry_run:
                for c in claimed:
                    number = c.get('RequestNumber', '?')
                    title = f"⏰ ESCALATION richiesta Ricezione — {number}"
                    msg = (f"Richiesta {number} ({_type_label(c.get('RequestType'))}) "
                           f"in attesa da oltre {ESCALATION_AGE_MINUTES} min senza risposta. "
                           f"Sollecitare il ricevimento.")
                    queue_popup(cur, 'INCOMING_RECEIVER', title, msg,
                                order_number=number, category='INCOMING')
            db.conn.commit()
    except Exception as e:
        logger.error("Escalation: claim/popup falliti: %s", e, exc_info=True)
        return 0

    if claimed and not dry_run:
        send_escalation_email(db, claimed)
        logger.info("Escalation: %d richieste notificate", len(claimed))
    elif claimed:
        logger.info("Escalation DRY-RUN: %d richieste eleggibili", len(claimed))
    return len(claimed)


def send_escalation_email(db, requests: list) -> None:
    """Email riepilogativa di escalation (una per tipo, a seconda dei destinatari)."""
    by_type = {}
    for r in requests:
        by_type.setdefault(r.get('RequestType'), []).append(r)
    for request_type, rows in by_type.items():
        recipients = _recipients_for_type(db, request_type)
        if not recipients:
            continue
        lines = [f"- {r.get('RequestNumber', '?')} | fornitore {r.get('SupplierName', '-')} "
                 f"| richiesta il {_fmt_dt(r.get('RequestedOn'))}"
                 for r in rows]
        subject = (f"[ESCALATION] {len(rows)} richieste Ricezione senza risposta "
                   f"— {_type_label(request_type)}")
        body = (f"Le seguenti richieste incoming ({_type_label(request_type)}) sono "
                f"in attesa da oltre {ESCALATION_AGE_MINUTES} minuti senza risposta:\n\n"
                + '\n'.join(lines)
                + "\n\nAccedi al modulo Ricezione (Soluzioni Incoming) per rispondere.")
        _send_email_async(recipients, subject, body)


# ───────────────────── Reminder periodici ─────────────────────────────── #

def check_and_send_reminders(db, dry_run: bool = False) -> int:
    """Sollecito per le richieste ancora PENDING.

    Per ogni tipo, la configurazione (get_email_config) definisce
    'reminders_per_day': massimo N solleciti al giorno per richiesta,
    distribuiti uniformemente nell'arco della giornata (il k-esimo
    sollecito, 0-based, e' consentito solo dopo k * 24/N ore da mezzanotte).
    Per ogni richiesta eleggibile: popup a 'INCOMING_RECEIVER'
    (category='INCOMING') + log su tabella reminder + UNA email riepilogativa
    per tipo ai destinatari configurati.

    Ritorna il numero di richieste sollecitate. Pensato per essere
    schedulato ogni 30-60 minuti (lock cross-PC sul job).
    """
    _ensure_job_row(db, JOB_REMINDERS)
    if dry_run:
        claimed_job = True
    else:
        claimed_job = claim_job_run(db, JOB_REMINDERS, lock_minutes=30)
        if not claimed_job:
            logger.info("Job %s non claimato: skip reminder", JOB_REMINDERS)
            return 0

    sent = 0
    try:
        try:
            pending = incoming_db.get_pending_for_reminders(db) or []
        except Exception as e:
            logger.error("Lettura richieste PENDING fallita: %s", e)
            return 0

        now = datetime.now()
        hours_since_midnight = now.hour + now.minute / 60.0
        eligible = []
        for req in pending:
            request_type = req.get('RequestType')
            try:
                cfg = incoming_db.get_email_config(db, request_type) or {}
            except Exception as e:
                logger.error("get_email_config(%s) fallita: %s", request_type, e)
                continue
            per_day = int(cfg.get('reminders_per_day') or 0)
            if per_day <= 0:
                continue
            already = incoming_db.count_reminders_today(db, req.get('Id'))
            if already >= per_day:
                continue
            # k-esimo sollecito del giorno (0-based) solo dopo k * 24/N ore
            if hours_since_midnight < already * (24.0 / per_day):
                continue
            eligible.append((req, cfg))

        if not eligible:
            if not dry_run:
                log_job_run(db, JOB_REMINDERS, 'OK', 'Nessuna richiesta da sollecitare')
            return 0

        by_type = {}
        for req, cfg in eligible:
            by_type.setdefault(req.get('RequestType'), []).append(req)

        if not dry_run:
            with db._lock:
                cur = _cursor(db)
                for request_type, rows in by_type.items():
                    for req in rows:
                        number = req.get('RequestNumber', '?')
                        age_min = 0
                        try:
                            requested_on = req.get('RequestedOn')
                            if requested_on:
                                age_min = int((now - requested_on).total_seconds() // 60)
                        except Exception:
                            pass
                        title = f"⏰ REMINDER richiesta Ricezione — {number}"
                        msg = (f"Richiesta {number} ({_type_label(request_type)}) "
                               f"ancora in attesa da {age_min} min. "
                               f"Sollecitare il ricevimento.")
                        queue_popup(cur, 'INCOMING_RECEIVER', title, msg,
                                    order_number=number, category='INCOMING')
                        incoming_db.log_reminder_sent(db, req.get('Id'))
                        sent += 1
                db.conn.commit()

        # Email riepilogativa per tipo
        for request_type, rows in by_type.items():
            if dry_run:
                continue
            recipients = _recipients_for_type(db, request_type)
            if not recipients:
                continue
            lines = [f"- {r.get('RequestNumber', '?')} | fornitore {r.get('SupplierName', '-')} "
                     f"| richiesta il {_fmt_dt(r.get('RequestedOn'))}"
                     for r in rows]
            subject = (f"[REMINDER] {len(rows)} richieste Ricezione in attesa "
                       f"— {_type_label(request_type)}")
            body = (f"Le seguenti richieste incoming ({_type_label(request_type)}) "
                    f"sono ancora in attesa di risposta:\n\n"
                    + '\n'.join(lines)
                    + "\n\nAccedi al modulo Ricezione (Soluzioni Incoming) per rispondere.")
            _send_email_async(recipients, subject, body)

        if not dry_run:
            log_job_run(db, JOB_REMINDERS, 'OK',
                        f"{sent} richieste sollecitate")
        logger.info("Reminder incoming: %d richieste sollecitate%s",
                    sent, ' (DRY-RUN)' if dry_run else '')
        return sent
    except Exception as e:
        logger.error("check_and_send_reminders fallito: %s", e, exc_info=True)
        if not dry_run:
            log_job_run(db, JOB_REMINDERS, 'ERROR', str(e)[:450])
            if claimed_job:
                release_job_lock(db, JOB_REMINDERS)
        return sent


# ───────────────────── Report mensile Excel ───────────────────────────── #

def _write_stats_sheet(ws, by_type: dict, title: str):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))

    columns = ('Tipo richiesta', 'Richieste', 'Risposte', 'Confermate OK',
               'Confermate KO', 'Tempo medio risposta (min)')
    ws.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border

    for request_type in incoming_db.REQUEST_TYPES:
        stats = (by_type or {}).get(request_type) or {}
        avg = stats.get('avg_response_minutes')
        avg_txt = round(float(avg), 1) if avg is not None else '-'
        row = (_type_label(request_type),
               stats.get('total', 0), stats.get('answered', 0),
               stats.get('confirmed_ok', 0), stats.get('confirmed_ko', 0),
               avg_txt)
        ws.append(row)
        for col_idx in range(1, len(columns) + 1):
            ws.cell(row=ws.max_row, column=col_idx).border = border

    widths = (28, 12, 12, 16, 16, 26)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{chr(64 + len(columns))}{ws.max_row}"


def generate_monthly_report_excel(stats: dict, year: int, month: int,
                                  file_path: str) -> str:
    """Genera il report Excel mensile (foglio 'Mese' + foglio 'YTD').
    Ritorna il percorso del file, None in caso di errore."""
    import openpyxl

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mese"
        _write_stats_sheet(ws, (stats.get('month') or {}).get('by_type'),
                           f"Mese {month:02d}/{year}")
        ws_ytd = wb.create_sheet("YTD")
        _write_stats_sheet(ws_ytd, (stats.get('ytd') or {}).get('by_type'),
                           f"YTD {year}")
        wb.save(file_path)
        logger.info("Report mensile incoming generato: %s", file_path)
        return file_path
    except Exception as e:
        logger.error("Generazione report Excel incoming fallita: %s", e, exc_info=True)
        return None


def generate_and_send_monthly_report(db, year: int = None, month: int = None,
                                     force: bool = False,
                                     dry_run: bool = False) -> bool:
    """Genera e invia il report mensile richieste incoming (Excel in allegato).

    Default: mese precedente. Dedup/lock cross-PC sul job
    'incoming_monthly_report' (claim prima di generare/inviare: un solo invio
    al mese anche con piu' PC). Destinatari: settings
    atribute='Incoming_soluzione_problemi' (incoming_db.get_monthly_recipients).

    Args:
        force: riesegui anche se il lock e' di un altro PC (run manuale).
        dry_run: genera l'Excel senza claim ne' invio email.

    Returns:
        True se inviato (o dry-run riuscito), False altrimenti.
    """
    _ensure_job_row(db, JOB_MONTHLY)

    today = datetime.now()
    if year is None or month is None:
        first_of_current = today.replace(day=1)
        prev = first_of_current.fromordinal(first_of_current.toordinal() - 1)
        year = year if year is not None else prev.year
        month = month if month is not None else prev.month

    claimed = False
    if not dry_run:
        claimed = (force_claim_job(db, JOB_MONTHLY, lock_minutes=60) if force
                   else claim_job_run(db, JOB_MONTHLY, lock_minutes=60))
        if not claimed:
            logger.info("Job %s non claimato: report mensile saltato", JOB_MONTHLY)
            return False

    excel_file = None
    try:
        stats = incoming_db.get_monthly_stats(db, year, month)
        month_name = datetime(year, month, 1).strftime('%B')
        excel_file = os.path.join(
            tempfile.gettempdir(), f"incoming_monthly_report_{year}{month:02d}.xlsx")
        if not generate_monthly_report_excel(stats, year, month, excel_file):
            raise RuntimeError("Generazione Excel fallita")

        if dry_run:
            log_job_run(db, JOB_MONTHLY, 'OK', f'DRY-RUN report {month:02d}/{year} generato')
            print(f"DRY-RUN: report {month:02d}/{year} generato in {excel_file}")
            return True

        recipients = incoming_db.get_monthly_recipients(db)
        if not recipients:
            logger.warning("Nessun destinatario (%s): report mensile NON inviato",
                           incoming_db.MONTHLY_RECIPIENTS_ATTRIBUTE)
            log_job_run(db, JOB_MONTHLY, 'SKIPPED',
                        f'Nessun destinatario per {incoming_db.MONTHLY_RECIPIENTS_ATTRIBUTE}')
            release_job_lock(db, JOB_MONTHLY)
            return False

        month_stats = (stats.get('month') or {}).get('by_type') or {}
        total = sum((s or {}).get('total', 0) for s in month_stats.values())
        answered = sum((s or {}).get('answered', 0) for s in month_stats.values())
        ko = sum((s or {}).get('confirmed_ko', 0) for s in month_stats.values())

        subject = f"Report mensile Ricezione (Incoming) — {month_name} {year}"
        body = (f"In allegato il report mensile delle richieste incoming per "
                f"{month_name} {year} (con riepilogo YTD).\n\n"
                f"Sintesi del mese: {total} richieste, {answered} risposte, "
                f"{ko} confermate KO.\n\n"
                f"Dettaglio per tipo di richiesta nel file Excel.")
        try:
            utils.send_email(recipients, subject, body, attachments=[excel_file])
            email_sent = True
        except Exception as e:
            logger.error("Invio report mensile incoming fallito: %s", e, exc_info=True)
            email_sent = False

        if email_sent:
            log_job_run(db, JOB_MONTHLY, 'OK',
                        f"Report {month:02d}/{year} inviato a {len(recipients)} destinatari")
            logger.info("Report mensile incoming %02d/%d inviato", month, year)
            return True
        log_job_run(db, JOB_MONTHLY, 'ERROR', 'Invio email fallito')
        release_job_lock(db, JOB_MONTHLY)
        return False
    except Exception as e:
        logger.error("generate_and_send_monthly_report fallito: %s", e, exc_info=True)
        log_job_run(db, JOB_MONTHLY, 'ERROR', str(e)[:450])
        if not dry_run and claimed:
            release_job_lock(db, JOB_MONTHLY)
        return False
    finally:
        if excel_file and os.path.exists(excel_file) and not dry_run:
            try:
                os.remove(excel_file)
            except OSError:
                pass


# ───────────────────── Standalone (Task Scheduler) ────────────────────── #

class _StandaloneDB:
    """Shim minimo per usare le funzioni incoming_db con una connessione
    pyodbc grezza (stessa interfaccia usata dalla classe Database di main.py:
    _lock, conn, cursor, _ensure_connection)."""

    def __init__(self, conn):
        self.conn = conn
        self._lock = threading.RLock()

    @property
    def cursor(self):
        return self.conn.cursor()

    def _ensure_connection(self):
        try:
            self.conn.cursor().execute("SELECT 1")
        except Exception:
            self.conn = _raw_connection()


def _raw_connection():
    import pyodbc
    from config_manager import ConfigManager
    cfg = ConfigManager(key_file='encryption_key.key',
                        config_file='db_config.enc').load_config()
    return pyodbc.connect(
        f"DRIVER={cfg['driver']};SERVER={cfg['server']};DATABASE={cfg['database']};"
        f"UID={cfg['username']};PWD={cfg['password']};MARS_Connection=Yes;"
        f"TrustServerCertificate=Yes"
    )


def get_standalone_db() -> _StandaloneDB:
    return _StandaloneDB(_raw_connection())


def main():
    import argparse
    parser = argparse.ArgumentParser(description='Job email modulo Ricezione (Incoming)')
    sub = parser.add_subparsers(dest='command', required=True)
    p_rem = sub.add_parser('reminders', help='Sollecito richieste PENDING')
    p_rem.add_argument('--dry-run', action='store_true')
    p_mon = sub.add_parser('monthly', help='Report mensile Excel')
    p_mon.add_argument('--year', type=int, default=None)
    p_mon.add_argument('--month', type=int, default=None)
    p_mon.add_argument('--force', action='store_true')
    p_mon.add_argument('--dry-run', action='store_true')
    p_esc = sub.add_parser('escalations', help='Popup+email escalation PENDING >= 120 min')
    p_esc.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()

    db = get_standalone_db()
    try:
        if args.command == 'reminders':
            n = check_and_send_reminders(db, dry_run=args.dry_run)
            print(f"Reminder: {n} richieste sollecitate"
                  + (' (DRY-RUN)' if args.dry_run else ''))
        elif args.command == 'monthly':
            ok = generate_and_send_monthly_report(
                db, year=args.year, month=args.month,
                force=args.force, dry_run=args.dry_run)
            print(f"Report mensile: {'OK' if ok else 'NON inviato'}"
                  + (' (DRY-RUN)' if args.dry_run else ''))
        elif args.command == 'escalations':
            n = process_escalations(db, dry_run=args.dry_run)
            print(f"Escalation: {n} richieste notificate"
                  + (' (DRY-RUN)' if args.dry_run else ''))
    finally:
        try:
            db.conn.close()
        except Exception:
            pass


if __name__ == '__main__':
    main()
