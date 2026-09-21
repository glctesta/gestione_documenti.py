# -*- coding: utf-8 -*-
"""
excel_gen.py — Esportazione Excel del kanban (pattern pdf_gen.py).

Un foglio per area (WIP / REPAIR): tutte le posizioni esistenti (anche vuote)
con le schede attive contenute. Intestazioni formattate + filtri + freeze panes.
"""
import os
import tempfile
import logging
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("WipKanban")

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="BBBBBB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_DATETIME_FMT = "DD/MM/YYYY HH:MM"

# (chiave riga, chiave i18n)
_FIELDS = [
    ("PositionCode", "col_position"),
    ("LabelCode", "col_labelcode"),
    ("OrderNumber", "col_order"),
    ("ProductCode", "col_product"),
    ("PhaseName", "col_phase"),
    ("ScanResult", "col_result"),
    ("DateIn", "col_datein"),
    ("LoadedBy", "col_loadedby"),
    ("ScanTimeFinish", "col_last_scan"),
]


def _auto_width(ws):
    for idx, col in enumerate(ws.iter_cols(1, ws.max_column), start=1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 4, 10), 42)


def kanban_workbook(rows_by_area, ui, path=None):
    """Crea il workbook Excel e restituisce il path del file temporaneo.

    rows_by_area: {'WIP': [row, ...], 'REPAIR': [row, ...]} (row = dict).
    ui: dizionario traduzioni per le intestazioni."""
    if not path:
        path = os.path.join(tempfile.gettempdir(),
                            f"kanban_export_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb = Workbook()
    headers = [ui.get(label_key, label_key) for _, label_key in _FIELDS]
    first = True
    for area in ("WIP", "REPAIR"):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = ui.get(f"area_{area.lower()}", area)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="center")
        for row in rows_by_area.get(area, []):
            ws.append([row.get(key) for key, _ in _FIELDS])
        for r in ws.iter_rows(min_row=2):
            for cell in r:
                cell.border = _BORDER
                if isinstance(cell.value, datetime):
                    cell.number_format = _DATETIME_FMT
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
        _auto_width(ws)
    wb.save(path)
    return path


# (chiave riga, chiave i18n) — report schede attive (pagina /kanban/report)
_REPORT_FIELDS = [
    ("ProductCode", "col_product"),
    ("Area", "col_type"),
    ("PositionCode", "col_box"),
    ("Pallet", "col_pallet"),
    ("OrderNumber", "col_order"),
    ("LabelCode", "col_labelcode"),
    ("PhaseName", "col_phase"),
    ("ScanResult", "col_result"),
    ("DateIn", "col_loaded_at"),
    ("LoadedBy", "col_loaded_by"),
]


def report_workbook(rows, ui, path=None):
    """Workbook del report: un solo foglio, stesso stile di kanban_workbook.

    rows: lista di dict (gia' ordinate per raggruppamento prodotto/scatola/ordine).
    ui: dizionario traduzioni per le intestazioni."""
    if not path:
        path = os.path.join(tempfile.gettempdir(),
                            f"kanban_report_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = ui.get("report_sheet", "Kanban")
    headers = [ui.get(label_key, label_key) for _, label_key in _REPORT_FIELDS]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(key) for key, _ in _REPORT_FIELDS])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.border = _BORDER
            if isinstance(cell.value, datetime):
                cell.number_format = _DATETIME_FMT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    _auto_width(ws)
    wb.save(path)
    return path
