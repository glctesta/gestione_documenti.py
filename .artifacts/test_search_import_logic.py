# -*- coding: utf-8 -*-
"""Test logica nuova search + import su DB reale (rollback, nessuna scrittura)."""
import pyodbc
from config_manager import ConfigManager

creds = ConfigManager(key_file='encryption_key.key', config_file='db_config.enc').load_config()
conn_str = (f"DRIVER={creds['driver']};SERVER={creds['server']};DATABASE={creds['database']};"
            f"UID={creds['username']};PWD={creds['password']};TrustServerCertificate=Yes")
cnxn = pyodbc.connect(conn_str, autocommit=False)
cnxn.timeout = 300
cur = cnxn.cursor()

# ── 1) Search: match per codice prodotto → tutti gli ordini aperti ──
q = "PFZA+67A020102"  # codice prodotto visto nei test precedenti
cur.execute("""SET TRANSACTION ISOLATION LEVEL READ UNCOMMITTED;
               SELECT IDProduct FROM Traceability_RS.dbo.Products WHERE ProductCode LIKE ?
               SET TRANSACTION ISOLATION LEVEL READ COMMITTED;""", (f"%{q}%",))
pids = [r[0] for r in cur.fetchall()]
print("Prodotti matchati:", pids)
if pids:
    ph = ",".join("?" for _ in pids)
    cur.execute(f"""SET TRANSACTION ISOLATION LEVEL READ UNCOMMITTED;
        SELECT o.IDOrder, o.OrderNumber, o.IsFinished, o.OrderDate
        FROM Traceability_RS.dbo.Orders o
        WHERE o.IsFinished = 0 AND o.IDProduct IN ({ph})
        ORDER BY o.OrderDate DESC
        SET TRANSACTION ISOLATION LEVEL READ COMMITTED;""", pids)
    rows = cur.fetchall()
    print(f"Ordini APERTI del prodotto: {len(rows)}")
    for r in rows[:5]:
        print("  ", r)

# ── 2) Import: risoluzione codici + replace + insert (ROLLBACK) ──
import openpyxl, io
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["CodiceProdotto", "CodiceEtichetta", "QuantitaPerPezzo"])
ws.append(["PFZA+67A020102|RP", "99900393", 2])       # esistente
ws.append(["PFZA+67A020102|RP", "CODICE_INESISTENTE", 1])  # deve essere saltato con warning
ws.append(["PRODOTTO_NON_ESISTE", "99900393", 1])     # deve essere saltato con warning
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)

wb2 = openpyxl.load_workbook(buf, read_only=True, data_only=True)
rows = list(wb2.active.iter_rows(values_only=True))
entries = [(str(r[0]).strip(), str(r[1]).strip(), float(r[2])) for r in rows[1:] if r and r[0]]
print("Righe parsed:", entries)

products, labels = {}, {}
for pc, lc, _ in entries:
    if pc not in products:
        cur.execute("SELECT IDProduct FROM Traceability_RS.dbo.Products WHERE UPPER(ProductCode) = UPPER(?)", (pc,))
        r = cur.fetchone()
        products[pc] = r[0] if r else None
    if lc not in labels:
        cur.execute("""SELECT m.MaterialeId FROM Traceability_RS.ind.Materiali m
                       JOIN Traceability_RS.ind.FamigliaMateriali fm ON fm.FamigliaMaterialiId = m.FamigliaMaterialiId
                       WHERE UPPER(m.CodiceMateriale) = UPPER(?) AND fm.Famiglia = 'Labels'""", (lc,))
        r = cur.fetchone()
        labels[lc] = r[0] if r else None
print("Prodotti:", products)
print("Etichette:", labels)

pids_ok = sorted({p for p in products.values() if p})
for pid in pids_ok:
    cur.execute("""UPDATE Traceability_RS.ind.BomIndirectMaterials SET DateOut = GETDATE()
                   WHERE idProduct = ? AND DateOut IS NULL""", (pid,))
    print(f"Chiuse associazioni attive per prodotto {pid}")
ins = 0
for pc, lc, qty in entries:
    pid, mid = products[pc], labels[lc]
    if not pid or not mid:
        continue
    cur.execute("""INSERT INTO Traceability_RS.ind.BomIndirectMaterials
                   (idProduct, MaterialeId, DateIn, [User], QuantityPerPiece)
                   VALUES (?, ?, GETDATE(), 'DIAG', ?)""", (pid, mid, qty))
    ins += 1
print(f"Righe che sarebbero inserite: {ins}")

cur.execute("""SELECT COUNT(*) FROM Traceability_RS.ind.BomIndirectMaterials
               WHERE idProduct = ? AND DateOut IS NULL""", (pids_ok[0],))
print("Associazioni attive post-import (transazione):", cur.fetchone()[0])

cnxn.rollback()
print("ROLLBACK eseguito — nessuna modifica persistita.")
cur.close(); cnxn.close()
