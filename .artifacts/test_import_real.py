import io
import pyodbc
import openpyxl
from config_manager import ConfigManager

creds = ConfigManager(key_file='encryption_key.key', config_file='db_config.enc').load_config()
cs = (f"DRIVER={creds['driver']};SERVER={creds['server']};DATABASE={creds['database']};"
      f"UID={creds['username']};PWD={creds['password']};TrustServerCertificate=Yes")
cn = pyodbc.connect(cs, autocommit=False)
cn.timeout = 120
cur = cn.cursor()

wb = openpyxl.Workbook()
ws = wb.active
ws.append(["CodiceProdotto", "CodiceEtichetta", "QuantitaPerPezzo"])
ws.append(["PFZA+67A020102|RP", "99900019", 2])
ws.append(["PFZA+67A020102|RP", "99900047", 0.5])
ws.append(["PFZA+67A020102|RP", "CODICE_INESISTENTE", 1])
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)

wb2 = openpyxl.load_workbook(buf, read_only=True, data_only=True)
rows = list(wb2.active.iter_rows(values_only=True))
entries = [(str(r[0]).strip(), str(r[1]).strip(), float(r[2])) for r in rows[1:] if r and r[0]]

products, labels = {}, {}
for pc, lc, _ in entries:
    if pc not in products:
        cur.execute("SELECT IDProduct FROM Traceability_RS.dbo.Products WHERE UPPER(ProductCode) = UPPER(?)", (pc,))
        r = cur.fetchone()
        products[pc] = r[0] if r else None
    if lc not in labels:
        cur.execute("""SELECT m.MaterialeId FROM Traceability_RS.ind.Materiali m
                       JOIN Traceability_RS.ind.FamigliaMateriali fm ON fm.FamigliaMaterialiId = m.FamigliaMaterialiId
                       WHERE UPPER(m.CodiceMateriale) = UPPER(?) AND fm.Famiglia = 'Labels'""", (lc,))
        r = cur.fetchone()
        labels[lc] = r[0] if r else None

warnings = []
for pc, pid in products.items():
    if pid is None:
        warnings.append(f"{pc}: prodotto non trovato")
for lc, mid in labels.items():
    if mid is None:
        warnings.append(f"{lc}: etichetta non trovata (famiglia 'Labels')")

pids_ok = sorted({p for p in products.values() if p})
for pid in pids_ok:
    cur.execute("""UPDATE Traceability_RS.ind.BomIndirectMaterials SET DateOut = GETDATE()
                   WHERE idProduct = ? AND DateOut IS NULL""", (pid,))
ins = 0
for pc, lc, qty in entries:
    pid, mid = products[pc], labels[lc]
    if not pid or not mid:
        continue
    cur.execute("""INSERT INTO Traceability_RS.ind.BomIndirectMaterials
                   (idProduct, MaterialeId, DateIn, [User], QuantityPerPiece)
                   VALUES (?, ?, GETDATE(), 'DIAG', ?)""", (pid, mid, qty))
    ins += 1

print("Prodotti:", products)
print("Etichette:", labels)
print("Inserite:", ins)
print("Warnings:", warnings)

cur.execute("""SELECT m.CodiceMateriale, b.QuantityPerPiece, b.[User]
               FROM Traceability_RS.ind.BomIndirectMaterials b
               JOIN Traceability_RS.ind.Materiali m ON m.MaterialeId = b.MaterialeId
               WHERE b.idProduct = ? AND b.DateOut IS NULL""", (pids_ok[0],))
print("Stato attivo post-import:", cur.fetchall())

cn.rollback()
print("ROLLBACK — nessuna modifica persistita.")
cur.close()
cn.close()
